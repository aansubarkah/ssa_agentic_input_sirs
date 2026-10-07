"""Repair the sample Excel files so they follow the web algorithm.

The sandbox computes Pasien Akhir Bulan (M) and Jumlah Hari Perawatan
(N) itself and rejects a row when the computed akhir is negative. This
script makes both sample files consistent with that algorithm:

- rows 5-40 of every monthly sheet (YYYY-MM): when the computed akhir
  is negative, the deficit is added to column D (Pasien Masuk), the
  safest column to adjust because no other web rule involves it;
- M and N are rewritten as static values (not formulas) so data
  readers such as pandas and openpyxl data_only can verify them;
- the TOTAL row 41 is rewritten as the column sums of C..U.

Deterministic: the output depends only on the input file. Usage:

    python scripts/perbaiki_excel.py [file.xlsx ...]

Without arguments it repairs both sample files in the repository.
"""

import re
import sys
from pathlib import Path

import openpyxl

# Excel column numbers (1 based): C..K movement counts, L days of care,
# M computed akhir, N computed hari rawat, O..T class bed days, U beds.
MOVEMENT_COLUMNS = (3, 4, 5, 6, 7, 8, 9, 10, 11)  # C..K
CLASS_COLUMNS = (15, 16, 17, 18, 19, 20)          # O..T
COLUMN_M = 13
COLUMN_N = 14
COLUMN_D = 4
DATA_ROW_START = 5
DATA_ROW_END = 40
TOTAL_ROW = 41


def to_int(raw):
    """Same number coercion as sirs_lib.read_excel."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return 0
    try:
        return max(int(float(raw)), 0)
    except (TypeError, ValueError):
        return 0


def computed_akhir(values):
    return (values[0] + values[1] + values[2]
            - (values[3] + values[4] + values[5]
               + values[6] + values[7] + values[8]))


def repair(path):
    """Repair one workbook in place. Returns a stats dict."""
    workbook = openpyxl.load_workbook(path)
    stats = {"sheets": 0, "rows": 0, "masuk_adjusted": 0}
    for name in workbook.sheetnames:
        if not re.fullmatch(r"\d{4}-\d{2}", str(name).strip()):
            continue  # skip the Info sheet and other non monthly sheets
        sheet = workbook[name]
        totals = {col: 0 for col in range(3, 22)}
        for row_index in range(DATA_ROW_START, DATA_ROW_END + 1):
            movement = [to_int(sheet.cell(row_index, c).value)
                        for c in MOVEMENT_COLUMNS]
            akhir = computed_akhir(movement)
            if akhir < 0:
                deficit = -akhir
                sheet.cell(row_index, COLUMN_D).value = (
                    movement[1] + deficit)
                movement[1] += deficit
                akhir = 0
                stats["masuk_adjusted"] += 1
            hari = sum(to_int(sheet.cell(row_index, c).value)
                       for c in CLASS_COLUMNS)
            sheet.cell(row_index, COLUMN_M).value = akhir
            sheet.cell(row_index, COLUMN_N).value = hari
            for col in range(3, 22):
                totals[col] += to_int(sheet.cell(row_index, col).value)
            stats["rows"] += 1
        for col in range(3, 22):
            sheet.cell(TOTAL_ROW, col).value = totals[col]
        stats["sheets"] += 1
    workbook.save(path)
    return stats


def main():
    here = Path(__file__).resolve().parent.parent
    files = sys.argv[1:] or [
        here / "RL32-dummy-24-bulan.xlsx",
        here / "RL32-latihan-20-data.xlsx",
    ]
    for target in files:
        stats = repair(str(target))
        print("%s: %d sheet, %d baris, %d baris masuk ditambah defisit"
              % (target, stats["sheets"], stats["rows"],
                 stats["masuk_adjusted"]))


if __name__ == "__main__":
    main()
